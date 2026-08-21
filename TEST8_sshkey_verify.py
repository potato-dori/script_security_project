import time
from openpyxl import load_workbook


# 결과 파일
file_path = r"C:\Git\script_security_project\RESULT\securitytest_result.xlsx"
save_path = file_path

# 장비 정보
device = "SWITCH"
console = "/SERIAL COM6 /BAUD 115200"

# 관리자 계정
username_admin = "admin"
password_admin = "Changeme1357!"

# SDN 시험 계정
sdn_username = "sdnuser"
sdn_password = "Changeme1357#"

# SFTP 서버
sftp_ip = "10.100.249.132"
sftp_port = "22222"
sftp_username = "tester"
sftp_password = "password"

# 장비 관리 IP
device_ip = "10.100.54.12/16"


###################################
# Excel
###################################

wb = load_workbook(file_path)
sheet1 = wb["RESULT"]
sheet2 = wb["LOG"]


def get_next_row(sheet):

    for row in range(sheet.max_row, 0, -1):
        if any(cell.value for cell in sheet[row]):
            return row + 1

    return 1


###################################
# 전체 화면 읽기
###################################

all_lines = []


def read_all():

    global all_lines
    all_lines = []

    num_rows = 500
    num_cols = 300

    for row in range(1, num_rows + 1):
        line = crt.Screen.Get(row, 1, row, num_cols).rstrip()

        if line.strip() != "":
            all_lines.append(line)


###################################
# 시험 구간 추출
###################################

def select(all_lines, TEST_start, TEST_end):

    final_line = []
    capture = False

    for line in all_lines:

        if TEST_start in line:
            capture = True
            final_line.append(line)
            continue

        if capture:
            final_line.append(line)

            if TEST_end in line:
                break

    final = "\n".join(final_line)

    return final


###################################
# 결과 저장
###################################

def save_result(test_name, result, judge_value, final):

    row1 = get_next_row(sheet1)
    row2 = get_next_row(sheet2)

    sheet1.cell(row=row1, column=2, value=test_name)
    sheet1.cell(row=row1, column=3, value=result)
    sheet1.cell(row=row1, column=4, value=judge_value)

    sheet2.cell(row=row2, column=2, value=test_name)
    sheet2.cell(row=row2, column=3, value=final)

    wb.save(save_path)


###################################
# 장비 접속
###################################

def start():

    crt.Session.Connect(console)

    time.sleep(2)

    crt.Screen.Send("\n")
    crt.Screen.WaitForString("login")

    crt.Screen.Send(f"{username_admin}\r")
    crt.Screen.WaitForString("Password")

    crt.Screen.Send(f"{password_admin}\r")
    crt.Screen.WaitForString(f"{device}>")

    crt.Screen.Send("enable\n")
    crt.Screen.WaitForString(f"{device}#")


###################################
# config mode
###################################

def config_mode():

    crt.Screen.Send("con t\n")
    crt.Screen.WaitForString(f"{device}(config)#")


##########################################################################
# Port 1 Link Up 확인
##########################################################################

def check_port1_link():

    crt.Screen.Send("PORT1_LINK_CHECK_start\n")

    crt.Screen.Send("show port\n")
    crt.Screen.WaitForString(f"{device}#")

    crt.Screen.Send("PORT1_LINK_CHECK_end\n")

    time.sleep(1)

    read_all()

    final = select(
        all_lines,
        "PORT1_LINK_CHECK_start",
        "PORT1_LINK_CHECK_end"
    )

    for line in final.splitlines():

        parts = line.split()

        if len(parts) >= 3:

            if parts[0] == "1" and parts[2] == "up":
                print ("Port 1 link up")
                return True
    
    return False


###################################
# SFTP 서버 Ping 확인
###################################

def check_sftp_ping():

    crt.Screen.Send("SFTP_PING_CHECK_start\n")
    crt.Screen.Send(f"ping {sftp_ip}\n")

    time.sleep(3)

    # Ctrl + C
    crt.Screen.Send("\x03")
    crt.Screen.WaitForString(f"{device}#")
    crt.Screen.Send("SFTP_PING_CHECK_end\n")

    time.sleep(1)

    read_all()

    final = select(
        all_lines,
        "SFTP_PING_CHECK_start",
        "SFTP_PING_CHECK_end"
    )

    for line in final.splitlines():

        if "packets received" in line:

            parts = line.split(",")

            if len(parts) >= 2:

                received = parts[1].strip().split()[0]

                if int(received) > 0:
                    print("SFTP Server Ping Success")
                    return True

    print("SFTP Server Ping Fail")
    return False



##########################################################################
# 시험 환경 설정 (mgmt ip 설정)
##########################################################################

def setup_sdn():

    ###################################
    # Port 1 설정
    ###################################

    config_mode()

    crt.Screen.Send("port 1\n")
    crt.Screen.WaitForString(f"{device}(port[1])#")

    crt.Screen.Send("speed 1000\n")
    crt.Screen.WaitForString(f"{device}(port[1])#")

    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")

    # port link up 될 시간 확보
    time.sleep(5)

    ###################################
    # Port 1 Link Up 확인
    ###################################

    port1_up = check_port1_link()

    if port1_up == False:
        raise Exception("Port 1 Link Down or unknown")


    ###################################
    # VLAN 160 설정
    ###################################

    config_mode()

    crt.Screen.Send("vlan 160\n")
    crt.Screen.WaitForString(f"{device}(vlan-160)#")

    crt.Screen.Send("untagged 1\n")
    crt.Screen.WaitForString(f"{device}(vlan-160)#")

    crt.Screen.Send(f"ip address {device_ip}\n")
    crt.Screen.WaitForString(f"{device}(vlan-160)#")

    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")

    # 설정 후 통신 될 시간 필요
    time.sleep(5)


    ###################################
    # SFTP 서버 통신 확인
    ###################################

    ping_success = check_sftp_ping()

    if ping_success == False:
        raise Exception("STFP Server Ping Fail")


    ###################################
    # sdnuser 생성
    ###################################

    config_mode()

    crt.Screen.Send(f"username {sdn_username} password admin\n")
    crt.Screen.WaitForString("Password")

    crt.Screen.Send(f"{sdn_password}\r")
    crt.Screen.WaitForString("Please enter it again")

    crt.Screen.Send(f"{sdn_password}\r")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")


##########################################################################
# sdnuser 로그인
##########################################################################

def login_sdnuser():

    crt.Screen.Send("exit\n")
    crt.Screen.WaitForString("login")

    crt.Screen.Send(f"{sdn_username}\r")
    crt.Screen.WaitForString("Password")

    crt.Screen.Send(f"{sdn_password}\r")
    crt.Screen.WaitForString(f"{device}>")

    crt.Screen.Send("enable\n")
    crt.Screen.WaitForString(f"{device}#")


##########################################################################
# TEST8-02
# SSH 개인키 저장
##########################################################################

def TEST8_SSH_key():

    crt.Screen.Send(f"{test_name}_start\n")

    time.sleep(1)


    ###################################
    # 전송 전 sdnkey 확인
    ###################################

    crt.Screen.Send("show files\n")
    crt.Screen.WaitForString(f"{device}#")


    ###################################
    # SSH Key SFTP 전송
    ###################################

    crt.Screen.Send(
        f"cp sftp {sftp_ip} {sftp_port} "
        f"{sftp_username} {sftp_password} "
        f"trizkey.pub sdnkey\n"
    )

    crt.Screen.WaitForString("Success. Done.")
    crt.Screen.WaitForString(f"{device}#")


    ###################################
    # 전송 후 sdnkey 확인
    ###################################

    crt.Screen.Send("show files\n")
    crt.Screen.WaitForString(f"{device}#")

    crt.Screen.Send(f"{test_name}_end\n")

    time.sleep(1)


##########################################################################
# SSH 개인키 판정
##########################################################################

def TEST8_SSH_key_verify():

    read_all()

    final = select(
        all_lines,
        f"{test_name}_start",
        f"{test_name}_end"
    )

    success_count = 0
    key_count = 0

    for line in final.splitlines():

        if "Success. Done." in line:
            success_count += 1

        if "ssh_sdn_host_rsa_key" in line:
            key_count += 1

    if success_count >= 1 and key_count >= 1:
        result = "PASS"

    else:
        result = "FAIL"

    judge_value = f"SFTP Success={success_count}, Key File={key_count}"

    save_result(test_name, result, judge_value, final)


##########################################################################
# SDN Key 삭제
##########################################################################

def clear_sdnkey():

    crt.Screen.Send("clear sdnkey\n")
    crt.Screen.WaitForString(f"{device}#")


##########################################################################
# sdnuser 삭제
##########################################################################

def cleanup_sdnuser():

    ###################################
    # sdnuser 로그아웃
    ###################################

    crt.Screen.Send("exit\n")
    crt.Screen.WaitForString("login")


    ###################################
    # admin 로그인
    ###################################

    crt.Screen.Send(f"{username_admin}\r")
    crt.Screen.WaitForString("Password")

    crt.Screen.Send(f"{password_admin}\r")
    crt.Screen.WaitForString(f"{device}>")

    crt.Screen.Send("enable\n")
    crt.Screen.WaitForString(f"{device}#")


    ###################################
    # sdnuser 삭제
    ###################################

    config_mode()

    crt.Screen.Send(f"no username {sdn_username}\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")


##########################################################################
# Console 종료
##########################################################################

def disconnect():

    crt.Screen.Send("exit\n")
    crt.Screen.WaitForString("login")

    time.sleep(1)

    crt.Session.Disconnect()


##########################################################################
# 실행
##########################################################################

start()


##################################################
# SSH 개인키 시험 환경 구성
##################################################

setup_sdn()
login_sdnuser()


##################################################
# TEST8-02
# SSH 개인키 저장 확인
##################################################

test_name = "TEST8_SSH_private_key"

TEST8_SSH_key()
TEST8_SSH_key_verify()


##################################################
# SSH Key 삭제
##################################################

clear_sdnkey()


##################################################
# sdnuser 삭제
##################################################

cleanup_sdnuser()


##################################################
# Console 종료
##################################################

disconnect()


##################################################
# TEST8 완료 알림
##################################################

with open(
    r"C:\Git\script_security_project\TEST8_DONE.txt",
    "w"
) as f:
    f.write("DONE")





