import subprocess
from openpyxl import load_workbook

###################################
# 설정값 및 결과 파일
###################################

nmap_path = r"C:\Program Files (x86)\Nmap\nmap.exe"
target_ip = "10.100.54.12"
expected_file = r"C:\Git\script_security_project\TEST9_ssh_expected_algorithms.txt"

# 결과 파일
file_path = r"C:\Git\script_security_project\RESULT\securitytest_result.xlsx"
save_path = file_path

test_name = "TEST9_SSH_algorithms"


###################################
# Excel
###################################

wb = load_workbook(file_path)

sheet1 = wb["RESULT"]
sheet2 = wb["LOG"]


def get_next_row(sheet):

    for row in range(sheet.max_row, 0, -1):

        if any(cell.value for cell in sheet[row]):
            return row + 1

    return 1


###################################
# Excel 결과 저장
###################################

def save_result(result, judge_value, final):

    row1 = get_next_row(sheet1)
    row2 = get_next_row(sheet2)

    sheet1.cell(row=row1, column=2, value=test_name)
    sheet1.cell(row=row1, column=3, value=result)
    sheet1.cell(row=row1, column=4, value=judge_value)

    sheet2.cell(row=row2, column=2, value=test_name)
    sheet2.cell(row=row2, column=3, value=final)

    wb.save(save_path)


###################################
# 기준 알고리즘 파일 읽기
###################################

def load_expected_algorithms(file_path):

    algorithms = {}
    # 지금 읽고 있는 알고리즘이 어느 항목에 속하는지 기억하기 위한 변수
    current_section = None

    with open(file_path, "r", encoding="utf-8") as file:

        for line in file:

            # 문자열의 앞뒤 공백, 줄바꿈 제거
            line = line.strip()

            # 빈 줄은 무시
            if not line:
                continue

            # ":"로 끝나면 알고리즘 구분 항목
            if line.endswith(":"):

                # 마지막 문자 하나를 제외하고 가져오는 slicing
                current_section = line[:-1]
                # 딕셔너리 개념
                algorithms[current_section] = []


            # 알고리즘 값
            else:

                if current_section:
                    algorithms[current_section].append(line)


    return algorithms


###################################
# Nmap 실행
###################################

def run_nmap():

    result = subprocess.run(
        [
            nmap_path,
            "-sV",
            "-p", "22",
            "-n",
            "--script", "ssh2-enum-algos",
            target_ip
        ],
        # nmap 결과를 Python 변수에 저장
        capture_output=True,
        # nmap 결과를 문자열로 받기
        text=True
    )


    # Nmap 실행 자체가 실패한 경우
    if result.returncode != 0:

        print("Nmap 실행 실패")
        print(result.stderr)

        return None


    return result.stdout


##################################
# Nmap 결과에서 알고리즘 추출
###################################

def parse_nmap_algorithms(output):

    algorithms = {}

    target_sections = [
        "kex_algorithms",
        "server_host_key_algorithms",
        "encryption_algorithms",
        "mac_algorithms",
        "compression_algorithms"
    ]

    current_section = None


    for original_line in output.splitlines():

        ###################################
        # ssh2-enum-algos 출력이 아닌 경우
        ###################################

        if not original_line.strip().startswith("|"):

            if current_section:
                current_section = None

            continue


        ###################################
        # 앞의 "|", "_" 제거
        ###################################

        line = original_line.strip()
        line = line.lstrip("|_").strip()


        section_found = False


        ###################################
        # 알고리즘 구분 항목 확인
        ###################################

        for section in target_sections:

            if line.startswith(section + ":"):

                current_section = section
                algorithms[current_section] = []
                section_found = True

                break


        if section_found:

            continue


        ###################################
        # 알고리즘 값 저장
        ###################################

        if current_section:

            if line:
                algorithms[current_section].append(line)


    return algorithms


###################################
# 기준값과 실제값 비교
###################################

def compare_algorithms(expected, actual):

    pass_flag = True
    summary_lines = []
    pass_section_count = 0


    for section in expected:

        print()
        print("===================================")
        print(section)
        print("===================================")

        # 리스트를 집합으로 변경, 집한은 순서를 신경쓰지 않음
        expected_set = set(expected[section])
        actual_set = set(actual.get(section, []))


        ###################################
        # 누락 알고리즘 - 집합 빼기
        ###################################

        missing = expected_set - actual_set


        ###################################
        # 추가 알고리즘
        ###################################

        extra = actual_set - expected_set


        ###################################
        # 결과 판정
        ###################################

        #missing, extra 둘다 비어있으면
        if not missing and not extra:

            print("결과 : PASS")
            summary_lines.append("결과 : PASS")
            pass_section_count += 1


        else:

            print("결과 : FAIL")
            summary_lines.append("결과 : FAIL")
            pass_flag = False


            if missing: # missing이 true이면

                print()
                print("[누락된 알고리즘]")

                summary_lines.append("")
                summary_lines.append("[누락된 알고리즘]")

                for algorithm in sorted(missing):

                    print(f" - {algorithm}")
                    summary_lines.append(f" - {algorithm}")


            if extra: # extra가 true이면

                print()

                print("[추가로 확인된 알고리즘]")

                summary_lines.append("")
                summary_lines.append("[추가로 확인된 알고리즘]")

                for algorithm in sorted(extra):

                    print(f" - {algorithm}")
                    summary_lines.append(f" - {algorithm}")

        summary_lines.append("")


    summary = "\n".join(summary_lines)
    judge_value = f"{pass_section_count}/{len(expected)} section PASS"


    return pass_flag, judge_value, summary


###################################
# Main
###################################

expected = load_expected_algorithms(expected_file)


nmap_output = run_nmap()


if nmap_output is None:

    print()
    print("최종 결과 : FAIL")

    save_result(
        "FAIL",
        "Nmap 실행 실패",
        "Nmap 실행 실패"
    )

else:

    ###################################
    # Nmap 원본 결과 출력
    ###################################

    print()
    print("===================================")
    print("Nmap Result")
    print("===================================")

    print(nmap_output)


    ###################################
    # Nmap 결과 Parsing
    ###################################

    actual = parse_nmap_algorithms(nmap_output)


    ###################################
    # 알고리즘 비교
    ###################################

    final_result, judge_value, summary = compare_algorithms(
        expected,
        actual
    )


    ###################################
    # 최종 판정
    ###################################

    print()
    print("===================================")

    if final_result:

        print("최종 결과 : PASS")
        result =  "PASS"

    else:

        print("최종 결과 : FAIL")
        result = "FAIL"

    print("===================================")


    ###################################
    # LOG 내용
    ###################################

    final_log = (
        "[Algorithm Compare Result]\n\n"
        + summary
        + "\n\n"
        + "[Nmap Original Result]\n\n"
        + nmap_output
    )


    ###################################
    # Excel 저장
    ###################################

    save_result(
        result,
        judge_value,
        final_log
    )


