from openpyxl import load_workbook


###################################
# 결과 파일
###################################

file_path = r"C:\Git\script_security_project\RESULT\securitytest_result.xlsx"


###################################
# 다음 기록 Row 확인
###################################

def get_next_row(sheet):

    for row in range(sheet.max_row, 0, -1):

        if any(cell.value for cell in sheet[row]):
            return row + 1

    return 1


###################################
# 전체 화면 읽기
###################################

def read_all(crt):

    all_lines = []
    num_rows=500
    num_cols=300


    for row in range(1, num_rows + 1):

        line = crt.Screen.Get(
            row,
            1,
            row,
            num_cols
        ).rstrip()


        if line.strip() != "":
            all_lines.append(line)


    return all_lines


###################################
# 시험 구간 추출
###################################

def select(all_lines,TEST_start,TEST_end):

    final_line = []

    capture = False


    for line in all_lines:

        if TEST_start in line:

            capture = True
            final_line.append(line)

            continue


        if capture:

            final_line.append(line)


            if TEST_end in line:
                break


    return "\n".join(final_line)


###################################
# 특정 시험 구간 읽기
###################################

def test_part(
    crt,
    TEST_start,
    TEST_end,
):

    all_lines = read_all(crt)

    final = select(
        all_lines,
        TEST_start,
        TEST_end
    )

    return final


###################################
# RESULT / LOG 저장
###################################

def save_result(
    test_name,
    test_result,
    judge_value,
    final
):

    wb = load_workbook(file_path)

    sheet1 = wb["RESULT"]
    sheet2 = wb["LOG"]


    row1 = get_next_row(sheet1)
    row2 = get_next_row(sheet2)


    ###################################
    # RESULT
    ###################################

    sheet1.cell(
        row=row1,
        column=2,
        value=test_name
    )

    sheet1.cell(
        row=row1,
        column=3,
        value=test_result
    )

    sheet1.cell(
        row=row1,
        column=4,
        value=judge_value
    )


    ###################################
    # LOG
    ###################################

    sheet2.cell(
        row=row2,
        column=2,
        value=test_name
    )

    sheet2.cell(
        row=row2,
        column=3,
        value=final
    )


    wb.save(file_path)