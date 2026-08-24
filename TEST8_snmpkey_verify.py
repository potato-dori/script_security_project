import time
from openpyxl import load_workbook

# 결과 파일
file_path = r"C:\Git\script_security_project\RESULT\securitytest_result.xlsx"
save_path = file_path

# 장비 정보
device = "SWITCH"
console = "/SERIAL COM6 /BAUD 115200"

# 기존 관리자
username_admin = "admin"
password_admin = "Changeme1357!"

# SNMP 시험
snmp_auth_password = "Changeme1357!"
snmp_priv_password = "Changeme1357#"


###################################
# Excel
###################################

wb = load_workbook(file_path)
sheet1 = wb["RESULT"]
sheet2 = wb["LOG"]


def get_next_row(sheet):

    for row in range(sheet.max_row, 0, -1):

        if any(cell.value for cell in sheet[row]):
            return row + 1

    return 1


###################################
# 전체 화면 읽기
###################################

all_lines = []


def read_all():

    global all_lines
    all_lines = []

    num_rows = 500
    num_cols = 300

    for row in range(1, num_rows + 1):
        line = crt.Screen.Get(row, 1, row, num_cols).rstrip()

        if line.strip() != "":
            all_lines.append(line)



###################################
# 시험 구간 추출
###################################

def select(all_lines, TEST_start, TEST_end):

    final_line = []
    capture = False

    for line in all_lines:

        if TEST_start in line:
            capture = True
            final_line.append(line)
            continue

        if capture:
            final_line.append(line)

            if TEST_end in line:
                break

    return "\n".join(final_line)

###################################
# 결과 저장
###################################

def save_result(test_name, result, judge_value, final):

    row1 = get_next_row(sheet1)
    row2 = get_next_row(sheet2)

    sheet1.cell(row=row1, column=2, value=test_name)
    sheet1.cell(row=row1, column=3, value=result)
    sheet1.cell(row=row1, column=4, value=judge_value)

    sheet2.cell(row=row2, column=2, value=test_name)
    sheet2.cell(row=row2, column=3, value=final)

    wb.save(save_path)


###################################
# 장비 접속
###################################

def start():

    crt.Session.Connect(console)

    time.sleep(2)

    crt.Screen.Send("\n")
    crt.Screen.WaitForString("login")

    crt.Screen.Send(f"{username_admin}\r")
    crt.Screen.WaitForString("Password")

    crt.Screen.Send(f"{password_admin}\r")
    crt.Screen.WaitForString(f"{device}>")

    crt.Screen.Send("enable\n")
    crt.Screen.WaitForString(f"{device}#")


###################################
# config mode
###################################

def config_mode():

    crt.Screen.Send("con t\n")
    crt.Screen.WaitForString(f"{device}(config)#")


##########################################################################
# TEST8-01
# SNMP 인증 및 암호화 비밀번호 저장
##########################################################################

def TEST8_SNMP_secret():

    ###################################
    # SNMP 설정
    ###################################

    config_mode()

    crt.Screen.Send("snmp enable\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("snmp group v3g v3 snmptest\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("snmp view v3v included .1\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("snmp access v3g v3 auth read v3v write v3v notify v3v\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send(
        f"snmp v3 user snmptest auth sha {snmp_auth_password} "
        f"priv aes {snmp_priv_password}\n"
    )
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")


    ###################################
    # 여기부터 비밀번호 노출 확인
    ###################################

    crt.Screen.Send(f"{test_name}_start\n")

    time.sleep(1)

    ###################################
    # show snmp
    ###################################

    crt.Screen.Send("show snmp\n")
    crt.Screen.WaitForString(f"{device}#")

    ###################################
    # running-configuration
    ###################################

    crt.Screen.Send("show running-configuration |include snmp\n")
    crt.Screen.WaitForString(f"{device}#")

    crt.Screen.Send(f"{test_name}_end\n")

    time.sleep(1)


###################################
# SNMP 비밀번호 노출 판정
###################################

def TEST8_SNMP_verify():

    read_all()

    final = select(
        all_lines,
        f"{test_name}_start",
        f"{test_name}_end"
    )

    auth_count = 0
    priv_count = 0

    for line in final.splitlines():

        if snmp_auth_password in line:
            auth_count += 1

        if snmp_priv_password in line:
            priv_count += 1


    if auth_count == 0 and priv_count == 0:
        result = "PASS"

    else:
        result = "FAIL"


    judge_value = f"Auth PW={auth_count}, Priv PW={priv_count}"

    save_result(test_name, result, judge_value, final)


##########################################################################
# snmp 삭제
##########################################################################

def clear_snmp():

    config_mode()

    crt.Screen.Send("no snmp v3 user snmptest\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("no snmp access v3g\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("no snmp view v3v included .1\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("no snmp group v3g v3 snmptest\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("snmp disable\n")
    crt.Screen.WaitForString(f"{device}(config)#")

    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")


##########################################################################
# Console 종료
##########################################################################

def disconnect():

    crt.Screen.Send("exit\n")
    crt.Screen.WaitForString("login")

    time.sleep(1)

    crt.Session.Disconnect()



##########################################################################
# 실행
##########################################################################

start()


##################################################
# TEST8-01
# SNMP 인증 및 암호화 비밀번호 출력 확인
##################################################

test_name = "TEST8_SNMPkey_verify"

TEST8_SNMP_secret()
TEST8_SNMP_verify()


##################################################
# SNMP 설정 삭제
##################################################

clear_snmp()


##################################################
# 장비 로그아웃 및 Console 연결 종료
##################################################

disconnect()