import time
from openpyxl import load_workbook

# 결과 파일
file_path = r"C:\Git\script_security_project\RESULT\securitytest_result.xlsx"
save_path = file_path

# 장비 정보
device = "SWITCH"
console = "/SERIAL COM6 /BAUD 115200"

# 기존 관리자 계정
username_admin = "admin"
password_admin = "Changeme1357!"

# TEST6 시험용 계정
username_admin2 = "admin2"
username_user = "user"

# 모든 시험용 비밀번호
test_password = "Changeme1357#"


###################################
# Excel
###################################

wb = load_workbook(file_path)
sheet1 = wb["RESULT"]
sheet2 = wb["LOG"]


def get_next_row(sheet):
    for row in range(sheet.max_row, 0, -1):
        if any(cell.value for cell in sheet[row]):
            return row + 1
    return 1


###################################
# 전체 화면 읽기
###################################

all_lines = []


def read_all():

    global all_lines
    all_lines = []

    num_rows = 300
    num_cols = 300

    for row in range(1, num_rows + 1):
        line = crt.Screen.Get(row, 1, row, num_cols).rstrip()

        if line.strip() != "":
            all_lines.append(line)


###################################
# 시험 구간 추출
###################################

def select(all_lines, TEST_start, TEST_end):

    final_line = []
    capture = False

    for line in all_lines:

        if TEST_start in line:
            capture = True
            final_line.append(line)
            continue

        if capture:
            final_line.append(line)

            if TEST_end in line:
                break

    final = "\n".join(final_line)

    return final


###################################
# 비밀번호 출력 여부 판단
###################################

def password_mask_verify():

    read_all()

    final = select(
        all_lines,
        f"{test_name}_start",
        f"{test_name}_end"
    )

    j = 0

    for line in final.splitlines():

        if test_password in line:
            j += 1

    # 입력한 비밀번호가 화면에 한 번도 출력되지 않아야 PASS
    result = "PASS" if j == 0 else "FAIL"

    row1 = get_next_row(sheet1)
    row2 = get_next_row(sheet2)

    sheet1.cell(row=row1, column=2, value=test_name)
    sheet1.cell(row=row1, column=3, value=result)
    sheet1.cell(row=row1, column=4, value=j)

    sheet2.cell(row=row2, column=2, value=test_name)
    sheet2.cell(row=row2, column=3, value=final)

    wb.save(save_path)


###################################
# 장비 접속
###################################

def start():

    crt.Session.Connect(console)

    time.sleep(2)

    crt.Screen.Send("\n")
    crt.Screen.WaitForString("login")

    crt.Screen.Send(f"{username_admin}\r")
    crt.Screen.WaitForString("Password")

    crt.Screen.Send(f"{password_admin}\r")
    crt.Screen.WaitForString(f"{device}>")

    crt.Screen.Send("enable\n")
    crt.Screen.WaitForString(f"{device}#")


###################################
# config mode 진입
###################################

def config_mode():

    crt.Screen.Send("conf t\n")
    crt.Screen.WaitForString(f"{device}(config)# ")


##########################################################################
# TEST6 : 1.4.1 비밀번호 출력
##########################################################################


##################################################
# TEST6-01
# 관리자 계정 비밀번호 설정 시 비밀번호 출력
##################################################

def TEST6_admin_password_set():
    crt.Screen.Send(f"{test_name}_start\n")
    time.sleep(1)
    config_mode()

    crt.Screen.Send(
        f"username {username_admin2} password admin\n"
    )
    crt.Screen.WaitForString("Password")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString("Please enter it again")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString(f"{device}(config)# ")
    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")

    crt.Screen.Send(f"{test_name}_end\n")
    time.sleep(1)


##################################################
# TEST6-02
# 사용자 계정 비밀번호 설정 시 비밀번호 출력
##################################################

def TEST6_user_password_set():
    crt.Screen.Send(f"{test_name}_start\n")
    time.sleep(1)

    config_mode()

    crt.Screen.Send(
        f"username {username_user} password guest\n"
    )
    crt.Screen.WaitForString("Password")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString("Please enter it again")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString(f"{device}(config)# ")
    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")
    crt.Screen.Send(f"{test_name}_end\n")
    time.sleep(1)

##################################################
# TEST6-03
# 운영모드 변경용 비밀번호 설정 시 비밀번호 출력
##################################################

def TEST6_enable_password_set():
    crt.Screen.Send(f"{test_name}_start\n")
    time.sleep(1)
    config_mode()

    crt.Screen.Send(
        f"enable {username_admin2} password\n"
    )
    crt.Screen.WaitForString("Enable Password")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString("Please enter it again")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString(f"{device}(config)# ")
    crt.Screen.Send("end\n")
    crt.Screen.WaitForString(f"{device}#")
    crt.Screen.Send(f"{test_name}_end\n")
    time.sleep(1)


##################################################
# TEST6-04, -05
# 관리자 계정 로그인 시 비밀번호 출력
# 운용모드 계정 로그인 시 비밀번호 출력
##################################################

def TEST6_admin_login():
    crt.Screen.Send(f"{test_name}_start\n")
    time.sleep(1)

    # 로그인 화면으로 이동
    crt.Screen.Send("exit\n")
    crt.Screen.WaitForString("login")
    crt.Screen.Send(f"{username_admin2}\r")
    crt.Screen.WaitForString("Password")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString(f"{device}>")

    #운용모드 확인
    crt.Screen.Send("enable\n")
    crt.Screen.WaitForString("Enable Password")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString(f"{device}#")
    
    crt.Screen.Send(f"{test_name}_end\n")
    time.sleep(1)


##################################################
# TEST6-06
# 사용자 계정 로그인 시 비밀번호 출력
##################################################

def TEST6_user_login():
    crt.Screen.Send(f"{test_name}_start\n")

    time.sleep(1)

    # 현재 admin2 enable 상태이므로 로그인 화면 이동
    crt.Screen.Send("exit\n")
    crt.Screen.WaitForString("login")
    crt.Screen.Send(f"{username_user}\r")
    crt.Screen.WaitForString("Password")
    crt.Screen.Send(f"{test_password}\r")
    crt.Screen.WaitForString(f"{device}>")

    crt.Screen.Send(f"{test_name}_end\n")
    time.sleep(1)


##########################################################################
# 실행
##########################################################################

start()

##################################################
# 관리자 계정 비밀번호 설정
##################################################

test_name = "TEST6_PW_mask_admin2_set"
TEST6_admin_password_set()
password_mask_verify()

##################################################
# 사용자 계정 비밀번호 설정
##################################################

test_name = "TEST6_PW_mask_user_set"
TEST6_user_password_set()
password_mask_verify()


##################################################
# 운영모드 변경용 비밀번호 설정
##################################################

test_name = "TEST6_PW_mask_enable_set"
TEST6_enable_password_set()
password_mask_verify()


##################################################
# 관리자 계정 로그인 # 운용모드 변경
##################################################
test_name = "TEST6_PW_mask_admin2&enable_login_"

TEST6_admin_login()
password_mask_verify()


##################################################
# 사용자 계정 로그인
##################################################

test_name = "TEST6_PW_mask_user_login"
TEST6_user_login()
password_mask_verify()
